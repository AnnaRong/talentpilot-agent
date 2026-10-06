"""Consent-bound overseas profile collection and resume export."""
from __future__ import annotations

import hashlib
import io
import json
import re
import time
import unicodedata
import uuid
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import unquote, urlsplit, urlunsplit

from docx import Document

from .config import ROOT, settings
from .platform_store import digest, private_write, utcnow

PROFILE_LIST_FIELDS = (
    "experience", "education", "skills", "languages", "certifications", "projects"
)
PROFILE_SECTIONS = (
    ("experience", "工作经历"),
    ("projects", "项目经历"),
    ("education", "教育经历"),
    ("skills", "技能"),
    ("languages", "语言"),
    ("certifications", "证书"),
)


def canonicalize_linkedin_url(value: str) -> str:
    parsed = urlsplit(value.strip())
    parts = [part for part in parsed.path.split("/") if part]
    try:
        port = parsed.port
    except ValueError as error:
        raise ValueError("只接受 https://www.linkedin.com/in/* 个人资料页") from error
    if (
        parsed.scheme != "https"
        or (parsed.hostname or "").lower() != "www.linkedin.com"
        or port not in (None, 443)
        or parsed.username is not None
        or parsed.password is not None
        or len(parts) < 2
        or parts[0] != "in"
    ):
        raise ValueError("只接受 https://www.linkedin.com/in/* 个人资料页")
    vanity = unicodedata.normalize("NFKC", unquote(parts[1])).strip().lstrip("@")
    if (
        not vanity
        or len(vanity) > 100
        or not any(character.isalnum() for character in vanity)
        or not all(character.isalnum() or character in "-_." for character in vanity)
    ):
        raise ValueError("LinkedIn vanity 无效")
    return urlunsplit(("https", "www.linkedin.com", f"/in/{vanity}", "", ""))


def latest_positions(store) -> list[dict]:
    with store.transaction() as cx:
        row = cx.execute(
            "SELECT payload FROM events WHERE type='italent.inventory.collected' "
            "ORDER BY id DESC LIMIT 1"
        ).fetchone()
    if not row:
        return []
    payload = store.unseal(row["payload"])
    return [
        {"job_id": str(item["job_id"]), "title": str(item.get("title") or item["job_id"])}
        for item in payload.get("positions", [])
        if item.get("job_id")
    ]


def _profile_markdown(profile: dict) -> str:
    lines = [f"# {profile['name']}"]
    for value in (profile.get("headline"), profile.get("location"), profile.get("linkedin_url")):
        if value:
            lines.extend(("", str(value)))
    if profile.get("about"):
        lines.extend(("", "## 个人概况", "", str(profile["about"])))
    for key, title in PROFILE_SECTIONS:
        items = profile.get(key) or []
        if items:
            lines.extend(("", f"## {title}", ""))
            lines.extend(f"- {item}" for item in items)
    return "\n".join(lines).strip() + "\n"


def _validate_profile(profile: dict) -> dict:
    if not isinstance(profile, dict):
        raise ValueError("海外候选人档案格式不正确")
    cleaned = {
        "name": str(profile.get("name") or "").strip(),
        "headline": str(profile.get("headline") or "").strip(),
        "location": str(profile.get("location") or "").strip(),
        "about": str(profile.get("about") or "").strip(),
        "linkedin_url": canonicalize_linkedin_url(str(profile.get("linkedin_url") or "")),
    }
    if not cleaned["name"] or len(cleaned["name"]) > 500:
        raise ValueError("候选人姓名缺失或过长")
    limits = {"headline": 1000, "location": 500, "about": 10_000}
    if any(len(cleaned[key]) > limit for key, limit in limits.items()):
        raise ValueError("候选人档案字段过长")
    for key in PROFILE_LIST_FIELDS:
        value = profile.get(key) or []
        if not isinstance(value, list):
            raise ValueError(f"{key} 必须是列表")
        maximum = 100 if key == "skills" else 50
        item_limit = 200 if key in {"skills", "languages"} else 5000
        items = [str(item).strip() for item in value if str(item).strip()]
        if len(items) > maximum or any(len(item) > item_limit for item in items):
            raise ValueError(f"{key} 超出档案上限")
        cleaned[key] = list(dict.fromkeys(items))
    return cleaned


def collect_profile(store, actor: str, payload: dict) -> dict:
    job_id = str(payload.get("job_id") or "").strip()
    job_title = str(payload.get("job_title") or "").strip()
    region = str(payload.get("region") or "").strip()
    consent_evidence = str(payload.get("consent_evidence") or "").strip()
    retention_days = payload.get("retention_days")
    if not job_id or not job_title:
        raise ValueError("必须绑定招聘岗位")
    if not region or not consent_evidence:
        raise ValueError("必须记录地区和合法处理授权依据")
    if not isinstance(retention_days, int) or not 1 <= retention_days <= 366:
        raise ValueError("保留期限必须为 1 至 366 天")
    profile = _validate_profile(payload.get("profile"))
    extracted_at = str(payload.get("extracted_at") or "").strip()
    try:
        captured = datetime.fromisoformat(extracted_at.replace("Z", "+00:00"))
    except ValueError as error:
        raise ValueError("采集时间格式不正确") from error
    if captured.tzinfo is None:
        captured = captured.replace(tzinfo=timezone.utc)
    if captured.timestamp() > time.time() + 300:
        raise ValueError("采集时间不能晚于当前时间")

    canonical_url = profile["linkedin_url"]
    candidate_id = hashlib.sha256(canonical_url.casefold().encode()).hexdigest()
    content_hash = digest(profile)
    record_id = uuid.uuid4().hex
    expires = time.time() + retention_days * 86400
    record = {
        "record_id": record_id,
        "candidate_id": candidate_id,
        "job_id": job_id,
        "job_title": job_title,
        "source": "linkedin_public_profile",
        "profile": profile,
        "markdown": _profile_markdown(profile),
        "captured_at": captured.astimezone(timezone.utc).isoformat(timespec="seconds"),
        "consent": {
            "region": region,
            "evidence": consent_evidence,
            "policy_version": "overseas-profile-v1",
            "expires_at": datetime.fromtimestamp(expires, timezone.utc).isoformat(timespec="seconds"),
        },
    }
    with store.transaction() as cx:
        existing = cx.execute(
            "SELECT id FROM overseas_resumes WHERE job_id=? AND candidate_id=? "
            "AND content_hash=? AND expires>?",
            (job_id, candidate_id, content_hash, time.time()),
        ).fetchone()
        if existing:
            return {"record_id": existing["id"], "candidate_id": candidate_id, "duplicate": True}
        cx.execute(
            "INSERT INTO overseas_resumes VALUES (?,?,?,?,?,?,?,?,?)",
            (record_id, candidate_id, job_id, job_title, record["source"], content_hash,
             store.seal(record), expires, utcnow()),
        )
        store.audit(cx, actor, "overseas.profile.collect", digest({"job_id": job_id, "record_id": record_id}))
    return {"record_id": record_id, "candidate_id": candidate_id, "duplicate": False}


def list_profiles(store, job_id: str) -> list[dict]:
    with store.transaction() as cx:
        rows = cx.execute(
            "SELECT id,payload,created_at,expires FROM overseas_resumes "
            "WHERE job_id=? AND expires>? ORDER BY created_at DESC",
            (job_id, time.time()),
        ).fetchall()
    result = []
    for row in rows:
        record = store.unseal(row["payload"])
        profile = record["profile"]
        result.append({
            "record_id": row["id"],
            "candidate_id": record["candidate_id"],
            "name": profile["name"],
            "headline": profile.get("headline", ""),
            "location": profile.get("location", ""),
            "linkedin_url": profile["linkedin_url"],
            "captured_at": record["captured_at"],
            "expires_at": record["consent"]["expires_at"],
        })
    return result


def _safe_filename(value: str, fallback: str) -> str:
    cleaned = re.sub(r"[\\/:*?\"<>|\x00-\x1f]+", "_", value).strip(" ._")
    return cleaned[:80] or fallback


def _resume_docx(record: dict) -> bytes:
    profile = record["profile"]
    document = Document()
    document.add_heading(profile["name"], 0)
    for value in (profile.get("headline"), profile.get("location"), profile.get("linkedin_url")):
        if value:
            document.add_paragraph(value)
    if profile.get("about"):
        document.add_heading("个人概况", level=1)
        document.add_paragraph(profile["about"])
    for key, title in PROFILE_SECTIONS:
        items = profile.get(key) or []
        if not items:
            continue
        document.add_heading(title, level=1)
        for item in items:
            document.add_paragraph(item, style="List Bullet")
    document.add_paragraph(f"来源：{record['source']} · 采集时间：{record['captured_at']}")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def export_resumes(store, actor: str, job_id: str, record_ids: list[str]) -> dict:
    if not 1 <= len(record_ids) <= 100 or len(record_ids) != len(set(record_ids)):
        raise ValueError("请选择 1 至 100 份不重复的海外简历")
    placeholders = ",".join("?" for _ in record_ids)
    with store.transaction() as cx:
        rows = cx.execute(
            f"SELECT id,payload FROM overseas_resumes WHERE job_id=? AND expires>? "
            f"AND id IN ({placeholders})",
            (job_id, time.time(), *record_ids),
        ).fetchall()
    if len(rows) != len(record_ids):
        raise ValueError("所选简历已过期、缺失或不属于当前岗位")
    records = [store.unseal(row["payload"]) for row in rows]
    export_id = uuid.uuid4().hex
    directory = settings.exports_dir / "overseas"
    directory.mkdir(parents=True, exist_ok=True, mode=0o700)
    path = directory / f"{export_id}.zip"
    used_names: set[str] = set()
    manifest = {"schema": "talentpilot-overseas-resume-export/v1", "job_id": job_id,
                "exported_at": utcnow(), "resumes": []}
    with io.BytesIO() as archive_bytes:
        with zipfile.ZipFile(archive_bytes, "w", zipfile.ZIP_DEFLATED) as archive:
            for index, record in enumerate(records, 1):
                base = _safe_filename(record["profile"]["name"], f"candidate-{index}")
                unique = base
                suffix = 2
                while unique.casefold() in used_names:
                    unique = f"{base}-{suffix}"
                    suffix += 1
                used_names.add(unique.casefold())
                archive.writestr(f"resumes/{unique}.docx", _resume_docx(record))
                archive.writestr(f"resumes/{unique}.md", record["markdown"])
                manifest["resumes"].append({
                    "record_id": record["record_id"], "candidate_id": record["candidate_id"],
                    "name": record["profile"]["name"], "source": record["source"],
                    "captured_at": record["captured_at"],
                })
            archive.writestr("manifest.json", json.dumps(manifest, ensure_ascii=False, indent=2))
        private_write(path, archive_bytes.getvalue())
    with store.transaction() as cx:
        store.audit(cx, actor, "overseas.resume.export", digest({"job_id": job_id, "records": record_ids}))
    return {"export_id": export_id, "count": len(records), "path": path}


def export_path(export_id: str) -> Path:
    if not re.fullmatch(r"[0-9a-f]{32}", export_id):
        raise ValueError("下载编号无效")
    path = (settings.exports_dir / "overseas" / f"{export_id}.zip").resolve()
    root = (settings.exports_dir / "overseas").resolve()
    if not path.is_relative_to(root) or not path.is_file():
        raise FileNotFoundError(export_id)
    return path


def extension_archive() -> Path:
    source = ROOT / "overseas-extension"
    files = ("README.md", "manifest.json", "popup.html", "popup.css", "popup.js", "service_worker.js", "content_script.js")
    if not all((source / name).is_file() for name in files):
        raise FileNotFoundError("overseas extension source missing")
    target = settings.exports_dir / "TalentPilot-LinkedIn-Collector.zip"
    target.parent.mkdir(parents=True, exist_ok=True, mode=0o700)
    newest = max((source / name).stat().st_mtime for name in files)
    if target.exists() and target.stat().st_mtime >= newest:
        return target
    temporary = target.with_suffix(".tmp")
    with zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as archive:
        for name in files:
            archive.write(source / name, f"TalentPilot-LinkedIn-Collector/{name}")
    temporary.replace(target)
    return target
